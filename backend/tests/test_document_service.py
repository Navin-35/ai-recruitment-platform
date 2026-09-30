import io
import fitz
import docx
from app.services.document_service import document_service


def test_document_content_hash():
    text1 = "Python developer with FastAPI experience."
    text2 = "  python   developer with   FastAPI experience. \n "
    hash1 = document_service.compute_content_hash(text1)
    hash2 = document_service.compute_content_hash(text2)
    assert hash1 == hash2
    assert len(hash1) == 64


def test_extract_text_from_pdf():
    # Create an in-memory PDF with fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "John Doe\nSoftware Engineer\nSkills: Python, Docker, PostgreSQL")
    pdf_bytes = doc.tobytes()
    doc.close()

    text, pages = document_service.extract_text_from_pdf(pdf_bytes)
    assert "John Doe" in text
    assert "Python" in text
    assert len(pages) == 1
    assert pages[0]["page_number"] == 1


def test_extract_text_from_docx():
    doc = docx.Document()
    doc.add_heading("Jane Smith", 0)
    doc.add_paragraph("Summary: Experienced Cloud Architect")
    doc.add_paragraph("Skills: AWS, Kubernetes, Terraform")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Degree"
    table.cell(0, 1).text = "B.S. CS"
    table.cell(1, 0).text = "Year"
    table.cell(1, 1).text = "2021"

    f = io.BytesIO()
    doc.save(f)
    docx_bytes = f.getvalue()

    text, pages = document_service.extract_text_from_docx(docx_bytes)
    assert "Jane Smith" in text
    assert "Cloud Architect" in text
    assert "AWS" in text
    assert "B.S. CS" in text
    assert len(pages) == 1


def test_detect_sections():
    resume_text = """
Jane Doe
jane@example.com

Summary:
Passionate backend engineer with 4 years experience building distributed microservices.

Skills:
Python, FastAPI, Docker, PostgreSQL, Redis, Kubernetes

Experience:
Senior Software Engineer at Acme Corp (2022 - Present)
- Architected REST APIs handling 10M daily requests.
- Optimized PostgreSQL queries reducing latency by 45%.

Projects:
Recruitment Matcher:
- Built hybrid retrieval engine with vector embeddings.

Education:
B.S. Computer Science from Tech University, 2021

Certifications:
AWS Certified Solutions Architect
"""
    sections = document_service.detect_sections(resume_text)
    assert "summary" in sections
    assert "skills" in sections
    assert "experience" in sections
    assert "projects" in sections
    assert "education" in sections
    assert "certifications" in sections
    assert "Acme Corp" in sections["experience"]
    assert "Python" in sections["skills"]
